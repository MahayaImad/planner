"""
Tests de l'API : authentification, sécurité et génération asynchrone.

    python3 tests/test_api.py
"""

import os
import sys
import time
from pathlib import Path

RACINE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RACINE))
os.environ.setdefault("DEBUG", "true")
os.environ.setdefault("DATABASE_URL", f"sqlite:///{RACINE}/test_api.db")

BASE = Path(os.environ["DATABASE_URL"].replace("sqlite:///", ""))
if BASE.exists():
    BASE.unlink()

from fastapi.testclient import TestClient  # noqa: E402

from app.main import app  # noqa: E402
from app.services import auth as service_auth  # noqa: E402

client = TestClient(app)


# ══════════════════════════════════════════════════════════════════
#  Mots de passe et jetons
# ══════════════════════════════════════════════════════════════════

def test_mot_de_passe_hache_avec_scrypt():
    empreinte = service_auth.hacher_mot_de_passe("motdepasse-solide")
    assert empreinte.startswith("scrypt$")
    # Deux hachages du même mot de passe diffèrent : le sel est aléatoire.
    assert empreinte != service_auth.hacher_mot_de_passe("motdepasse-solide")
    assert service_auth.verifier_mot_de_passe("motdepasse-solide", empreinte)
    assert not service_auth.verifier_mot_de_passe("autre", empreinte)
    assert not service_auth.doit_etre_rehache(empreinte)


def test_empreinte_heritee_est_relue_puis_remplacee():
    """Les comptes créés avant la migration doivent rester utilisables."""
    try:
        from passlib.hash import sha256_crypt
    except ImportError:
        return
    ancienne = sha256_crypt.hash("ancien-secret")
    assert service_auth.verifier_mot_de_passe("ancien-secret", ancienne)
    assert service_auth.doit_etre_rehache(ancienne)


def test_jeton_altere_est_rejete():
    jeton = service_auth.creer_token(1, 2, "a@b.dz")
    assert service_auth.decoder_token(jeton) is not None
    assert service_auth.decoder_token(jeton[:-4] + "AAAA") is None
    assert service_auth.decoder_token("nimportequoi") is None


def test_algorithme_none_est_rejete():
    """Un jeton annonçant « alg: none » ne doit jamais passer."""
    import base64
    import json
    jeton = service_auth.creer_token(1, 2, "a@b.dz")
    entete = base64.urlsafe_b64encode(
        json.dumps({"alg": "none", "typ": "JWT"}).encode()).rstrip(b"=").decode()
    forge = f"{entete}.{jeton.split('.')[1]}."
    assert service_auth.decoder_token(forge) is None


def test_cle_secrete_obligatoire_hors_developpement():
    from app.config import CLE_PAR_DEFAUT, Settings
    reglages = Settings(SECRET_KEY=CLE_PAR_DEFAUT, DEBUG=False)
    try:
        reglages.valider()
        assert False, "une clé par défaut ne doit pas être acceptée"
    except RuntimeError as e:
        assert "SECRET_KEY" in str(e)


def test_joker_cors_refuse_hors_developpement():
    from app.config import Settings
    reglages = Settings(SECRET_KEY="x" * 40, CORS_ORIGINS="*", DEBUG=False)
    try:
        reglages.valider()
        assert False, "le joker CORS ne doit pas être accepté"
    except RuntimeError as e:
        assert "CORS" in str(e)


# ══════════════════════════════════════════════════════════════════
#  Parcours complet
# ══════════════════════════════════════════════════════════════════

# (nom, type de salle exigé, heures/semaine, blocs de 2 h)
PROGRAMME = [("Maths", None, 4, 1), ("Arabe", None, 4, 1),
             ("Français", None, 3, 0), ("Sport", "sport", 2, 1)]


def _etablissement(suffixe: str, nb_classes: int = 2):
    """Inscrit une école avec ses ressources ; renvoie (entêtes, ids)."""
    r = client.post("/auth/inscrire", json={
        "ecole": {"nom": f"CEM {suffixe}", "email": f"{suffixe}@test.dz"},
        "admin": {"nom": "A", "prenom": "B", "email": f"admin{suffixe}@test.dz",
                  "mot_de_passe": "motdepasse-solide"}})
    assert r.status_code == 201, r.text
    entetes = {"Authorization": f"Bearer {r.json()['access_token']}"}

    matieres = {nom: client.post(
        "/matieres/", json={"nom": nom, "type_salle_requis": salle},
        headers=entetes).json()["id"] for nom, salle, _, _ in PROGRAMME}

    for i in range(nb_classes + 1):
        client.post("/salles/", json={"nom": f"S{i + 1}", "capacite": 40,
                                      "type": "classique"}, headers=entetes)
    for i in range(max(1, nb_classes // 4)):
        client.post("/salles/", json={"nom": f"Stade {i + 1}", "capacite": 60,
                                      "type": "sport"}, headers=entetes)

    classes = [client.post("/classes/", json={"nom": f"1AM{i + 1}",
                                              "niveau": "moyen", "effectif": 30},
                           headers=entetes).json()["id"]
               for i in range(nb_classes)]

    # Un enseignant ne peut pas couvrir plus de 32 créneaux : dimensionner
    # le corps enseignant sur le volume horaire réel de chaque matière.
    profs = {}
    for nom, _, heures, _ in PROGRAMME:
        nombre = max(1, (heures * nb_classes + 19) // 20)
        profs[nom] = [client.post("/professeurs/", json={
            "nom": f"P{nom}{k}", "prenom": "X", "max_heures_consecutives": 4,
            "matieres_ids": [matieres[nom]]}, headers=entetes).json()["id"]
            for k in range(nombre)]

    edt = client.post("/emplois-du-temps/",
                      json={"nom": "Semaine type", "annee_scolaire": "2025-2026"},
                      headers=entetes).json()["id"]
    return entetes, matieres, classes, profs, edt


def _demande(matieres, classes, profs, **extra):
    cours = []
    for rang, classe in enumerate(classes):
        for nom, _, heures, doubles in PROGRAMME:
            enseignants = profs[nom]
            cours.append({
                "classe_id": classe,
                "matiere_id": matieres[nom],
                "professeur_id": enseignants[rang % len(enseignants)],
                "heures_par_semaine": heures,
                "nb_seances_doubles": doubles,
            })
    return {"cours_requis": cours, "limite_secondes": 20,
            "grille": {"fermetures": [[2, [4, 5, 6]]]}, **extra}


def _attendre(entetes, edt, tache_id, delai=90):
    limite = time.time() + delai
    while time.time() < limite:
        tache = client.get(f"/emplois-du-temps/{edt}/taches/{tache_id}",
                           headers=entetes).json()
        if tache["terminee"]:
            return tache
        time.sleep(0.5)
    raise AssertionError("la tâche ne s'est jamais terminée")


def test_routes_protegees_sans_jeton():
    for methode, url in (("get", "/emplois-du-temps/"), ("get", "/professeurs/"),
                         ("get", "/classes/"), ("get", "/salles/")):
        r = getattr(client, methode)(url)
        assert r.status_code in (401, 403), f"{url} → {r.status_code}"


def test_generation_asynchrone_complete():
    entetes, matieres, classes, profs, edt = _etablissement("async")
    demande = _demande(matieres, classes, profs)

    depart = time.time()
    r = client.post(f"/emplois-du-temps/{edt}/generer", json=demande, headers=entetes)
    assert r.status_code == 202, r.text
    # La mise en file doit rendre la main tout de suite.
    assert time.time() - depart < 5.0
    tache_id = r.json()["id"]

    tache = _attendre(entetes, edt, tache_id)
    assert tache["statut"] == "terminee", tache["message"]
    attendu = sum(h for _, _, h, _ in PROGRAMME) * len(classes)
    assert tache["resultat"]["lecons_planifiees"] == attendu
    assert tache["resultat"]["qualite"]["trous_classes"] == 0

    lecons = client.get(f"/emplois-du-temps/{edt}/lecons", headers=entetes).json()
    assert len(lecons) == attendu


def test_statistiques_decrivent_l_emploi_du_temps_enregistre():
    """
    Les statistiques se recalculent depuis les leçons en base : elles
    doivent retrouver le total d'heures du programme et concorder avec
    les indicateurs que la génération a rapportés.
    """
    entetes, matieres, classes, profs, edt = _etablissement("stats")
    demande = _demande(matieres, classes, profs)
    tache_id = client.post(f"/emplois-du-temps/{edt}/generer", json=demande,
                           headers=entetes).json()["id"]
    tache = _attendre(entetes, edt, tache_id)
    assert tache["statut"] == "terminee", tache["message"]

    r = client.get(f"/emplois-du-temps/{edt}/statistiques", headers=entetes)
    assert r.status_code == 200, r.text
    stats = r.json()

    attendu = sum(h for _, _, h, _ in PROGRAMME) * len(classes)
    assert stats["totaux"]["lecons"] == attendu
    assert stats["totaux"]["lecons_hors_grille"] == 0
    assert stats["totaux"]["classes"] == len(classes)

    # Chaque division suit tout son programme, sans trou.
    assert len(stats["classes"]) == len(classes)
    for division in stats["classes"]:
        assert division["heures"] == sum(h for _, _, h, _ in PROGRAMME)
    assert stats["resume"]["trous_classes"] == tache["resultat"]["qualite"][
        "trous_classes"]
    assert stats["resume"]["trous_professeurs"] == tache["resultat"]["qualite"][
        "trous_professeurs"]

    # Le service de chaque professeur, et la somme qui doit retomber sur
    # le volume total.
    assert sum(p["heures"] for p in stats["professeurs"]) == attendu
    for enseignant in stats["professeurs"]:
        assert enseignant["heures"] > 0
        assert enseignant["jours_presence"] >= 1
        # Les clés d'un objet JSON sont des chaînes : la charge « 4 h »
        # revient indexée par « "4" ».
        assert sum(int(v) * n for v, n in enseignant["charges"].items()) == \
            enseignant["heures"]

    # Taux d'occupation : jamais plus de divisions occupées que possible.
    for seance in stats["occupation_seances"]:
        assert 0 <= seance["occupe"] <= seance["possible"]


def test_statistiques_refusent_un_emploi_du_temps_d_une_autre_ecole():
    entetes_a, *_, edt_a = _etablissement("stats-a")
    entetes_b, *_ = _etablissement("stats-b")
    r = client.get(f"/emplois-du-temps/{edt_a}/statistiques", headers=entetes_b)
    assert r.status_code == 404


def test_classeur_aller_retour_complet():
    """
    Export puis ré-import : le classeur doit se relire lui-même sans
    rien dupliquer, et sans toucher aux identifiants — les emplois du
    temps existants y renvoient.
    """
    import io as _io
    from openpyxl import load_workbook

    entetes, matieres, classes, profs, _edt = _etablissement("classeur")
    avant_classes = client.get("/classes/", headers=entetes).json()
    avant_profs = client.get("/professeurs/", headers=entetes).json()

    export = client.get("/donnees/export.xlsx", headers=entetes)
    assert export.status_code == 200, export.text
    classeur = load_workbook(_io.BytesIO(export.content))
    assert "Lisez-moi" in classeur.sheetnames
    for onglet in ("Matieres", "Salles", "Enseignants", "Classes", "Programme"):
        assert onglet in classeur.sheetnames
    # Les divisions créées doivent figurer dans l'export.
    noms = {ligne[0] for ligne in classeur["Classes"].iter_rows(
        min_row=2, values_only=True) if ligne[0]}
    assert noms == {c["nom"] for c in avant_classes}

    r = client.post("/donnees/importer",
                    files={"fichier": ("donnees.xlsx", export.content)},
                    headers=entetes)
    assert r.status_code == 200, r.text
    rapport = r.json()
    assert rapport["valide"], rapport["erreurs"]
    assert rapport["applique"]
    # Tout existait déjà : rien de créé.
    assert rapport["crees"]["Classes"] == 0
    assert rapport["crees"]["Enseignants"] == 0
    assert rapport["modifies"]["Classes"] == len(avant_classes)

    apres_classes = client.get("/classes/", headers=entetes).json()
    apres_profs = client.get("/professeurs/", headers=entetes).json()
    assert [c["id"] for c in apres_classes] == [c["id"] for c in avant_classes]
    assert [p["id"] for p in apres_profs] == [p["id"] for p in avant_profs]


def test_classeur_cree_un_etablissement_depuis_le_modele():
    """Le gabarit vierge, rempli, doit suffire à créer tout un CEM."""
    import io as _io
    from openpyxl import load_workbook

    r = client.post("/auth/inscrire", json={
        "ecole": {"nom": "CEM vierge", "email": "vierge@test.dz"},
        "admin": {"nom": "A", "prenom": "B", "email": "vierge-admin@test.dz",
                  "mot_de_passe": "motdepasse-solide"}})
    entetes = {"Authorization": f"Bearer {r.json()['access_token']}"}

    modele = client.get("/donnees/modele.xlsx", headers=entetes)
    assert modele.status_code == 200
    classeur = load_workbook(_io.BytesIO(modele.content))
    # Un gabarit est vide : seuls les en-têtes.
    assert classeur["Classes"].max_row == 1

    classeur["Matieres"].append(["Mathématiques", 4, ""])
    classeur["Matieres"].append(["Éducation physique", 1, "sport"])
    classeur["Salles"].append(["S1", 40, "classique"])
    classeur["Salles"].append(["Stade", 60, "sport"])
    classeur["Enseignants"].append(
        ["Benali", "Karim", "k@cem.dz", "", "Mathématiques", 4, 6, 20, "oui"])
    classeur["Enseignants"].append(
        ["Hadj", "Amina", "", "", "Éducation physique", 4, 6, 20, "non"])
    classeur["Classes"].append(["1AM1", "moyen", 32, "S1", 6])
    classeur["Programme"].append(["1AM1", "Mathématiques", "Benali", 5, 1, 2, "", ""])
    classeur["Programme"].append(
        ["1AM1", "Éducation physique", "Hadj", 2, 1, 2, "", ""])
    tampon = _io.BytesIO()
    classeur.save(tampon)

    # L'aperçu valide sans rien écrire.
    apercu = client.post("/donnees/importer?apercu=true",
                         files={"fichier": ("plein.xlsx", tampon.getvalue())},
                         headers=entetes)
    assert apercu.status_code == 200, apercu.text
    assert apercu.json()["valide"], apercu.json()["erreurs"]
    assert apercu.json()["applique"] is False
    assert client.get("/classes/", headers=entetes).json() == []

    r = client.post("/donnees/importer",
                    files={"fichier": ("plein.xlsx", tampon.getvalue())},
                    headers=entetes)
    rapport = r.json()
    assert rapport["valide"], rapport["erreurs"]
    assert rapport["crees"] == {"Matieres": 2, "Salles": 2, "Enseignants": 2,
                                "Classes": 1, "Programme": 2}

    classes = client.get("/classes/", headers=entetes).json()
    assert [c["nom"] for c in classes] == ["1AM1"]
    assert classes[0]["salle_attitree_id"] is not None
    prof = next(p for p in client.get("/professeurs/", headers=entetes).json()
                if p["nom"] == "Hadj")
    assert prof["assure_permanences"] is False
    assert len(client.get("/programme/", headers=entetes).json()) == 2


def test_classeur_fautif_n_ecrit_rien():
    """Une erreur, même en dernière feuille, annule tout l'import."""
    import io as _io
    from openpyxl import load_workbook

    r = client.post("/auth/inscrire", json={
        "ecole": {"nom": "CEM fautif", "email": "fautif@test.dz"},
        "admin": {"nom": "A", "prenom": "B", "email": "fautif-admin@test.dz",
                  "mot_de_passe": "motdepasse-solide"}})
    entetes = {"Authorization": f"Bearer {r.json()['access_token']}"}

    modele = client.get("/donnees/modele.xlsx", headers=entetes)
    classeur = load_workbook(_io.BytesIO(modele.content))
    classeur["Matieres"].append(["Mathématiques", 4, ""])
    classeur["Salles"].append(["S1", 40, "classique"])
    classeur["Enseignants"].append(
        ["Benali", "Karim", "", "", "Mathématiques", 4, 6, 20, "oui"])
    classeur["Classes"].append(["1AM1", "moyen", 32, "S1", 6])
    # Le programme désigne une matière qui n'existe nulle part : la
    # feuille Matières a pourtant été écrite avant, elle doit être
    # annulée avec le reste.
    classeur["Programme"].append(["1AM1", "Astronomie", "Benali", 2, 0, 1, "", ""])
    tampon = _io.BytesIO()
    classeur.save(tampon)

    r = client.post("/donnees/importer",
                    files={"fichier": ("fautif.xlsx", tampon.getvalue())},
                    headers=entetes)
    rapport = r.json()
    assert not rapport["valide"]
    assert any("Astronomie" in e for e in rapport["erreurs"]), rapport["erreurs"]
    assert client.get("/matieres/", headers=entetes).json() == []
    assert client.get("/classes/", headers=entetes).json() == []


def test_classeur_signale_une_salle_inconnue():
    import io as _io
    from openpyxl import load_workbook

    r = client.post("/auth/inscrire", json={
        "ecole": {"nom": "CEM salle", "email": "salle@test.dz"},
        "admin": {"nom": "A", "prenom": "B", "email": "salle-admin@test.dz",
                  "mot_de_passe": "motdepasse-solide"}})
    entetes = {"Authorization": f"Bearer {r.json()['access_token']}"}
    classeur = load_workbook(_io.BytesIO(
        client.get("/donnees/modele.xlsx", headers=entetes).content))
    classeur["Classes"].append(["1AM1", "moyen", 32, "Salle fantôme", 6])
    tampon = _io.BytesIO()
    classeur.save(tampon)

    rapport = client.post("/donnees/importer",
                          files={"fichier": ("x.xlsx", tampon.getvalue())},
                          headers=entetes).json()
    assert not rapport["valide"]
    assert any("Salle fantôme" in e for e in rapport["erreurs"]), rapport["erreurs"]


def test_une_seule_generation_a_la_fois():
    entetes, matieres, classes, profs, edt = _etablissement("concurrent")
    demande = _demande(matieres, classes, profs, limite_secondes=30)
    premiere = client.post(f"/emplois-du-temps/{edt}/generer", json=demande,
                           headers=entetes)
    assert premiere.status_code == 202
    seconde = client.post(f"/emplois-du-temps/{edt}/generer", json=demande,
                          headers=entetes)
    assert seconde.status_code == 409
    client.delete(f"/emplois-du-temps/{edt}/taches/{premiere.json()['id']}",
                  headers=entetes)
    _attendre(entetes, edt, premiere.json()["id"])


def test_annulation_interrompt_le_calcul():
    # Une instance jouet atteint l'optimum en moins d'une seconde : il
    # faut un problème que le solveur explore encore pour que
    # l'annulation ait un sens.
    entetes, matieres, classes, profs, edt = _etablissement("annul", nb_classes=14)
    demande = _demande(matieres, classes, profs, limite_secondes=600)
    tache_id = client.post(f"/emplois-du-temps/{edt}/generer", json=demande,
                           headers=entetes).json()["id"]
    time.sleep(4)
    en_cours = client.get(f"/emplois-du-temps/{edt}/taches/{tache_id}",
                          headers=entetes).json()
    assert not en_cours["terminee"], "le solveur a fini avant l'annulation"

    depart = time.time()
    r = client.delete(f"/emplois-du-temps/{edt}/taches/{tache_id}", headers=entetes)
    assert r.status_code == 200, r.text
    tache = _attendre(entetes, edt, tache_id, delai=60)
    assert tache["statut"] == "annulee"
    # L'arrêt doit être immédiat, sans attendre la limite de 600 s.
    assert time.time() - depart < 45


def test_diagnostic_signale_les_donnees_incoherentes():
    entetes, matieres, classes, profs, edt = _etablissement("diag")
    demande = _demande(matieres, classes, profs)
    r = client.post(f"/emplois-du-temps/{edt}/diagnostic", json=demande,
                    headers=entetes)
    assert r.status_code == 200 and r.json()["realisable"]

    # Professeur affecté à une matière qu'il n'enseigne pas.
    mauvais = dict(demande, cours_requis=[{
        "classe_id": classes[0], "matiere_id": matieres["Maths"],
        "professeur_id": profs["Arabe"][0], "heures_par_semaine": 4}])
    r = client.post(f"/emplois-du-temps/{edt}/diagnostic", json=mauvais,
                    headers=entetes)
    assert not r.json()["realisable"]
    assert "n'enseigne pas" in r.json()["erreurs"][0]

    # La génération refuse la demande sans occuper le pool de calcul.
    r = client.post(f"/emplois-du-temps/{edt}/generer", json=mauvais, headers=entetes)
    assert r.status_code == 422
    assert r.json()["detail"]["erreurs"]


def test_cloisonnement_entre_etablissements():
    """Une école ne doit jamais voir l'emploi du temps d'une autre."""
    entetes_a, matieres, classes, profs, edt_a = _etablissement("ecoleA")
    entetes_b, *_ = _etablissement("ecoleB")

    assert client.get(f"/emplois-du-temps/{edt_a}/lecons",
                      headers=entetes_b).status_code == 404
    assert client.post(f"/emplois-du-temps/{edt_a}/generer",
                       json=_demande(matieres, classes, profs),
                       headers=entetes_b).status_code == 404
    # L'école B ne voit que le sien, jamais celui de l'école A.
    visibles = {e["id"] for e in client.get("/emplois-du-temps/",
                                            headers=entetes_b).json()}
    assert edt_a not in visibles


def test_limite_de_calcul_est_plafonnee():
    from app.config import settings
    entetes, matieres, classes, profs, edt = _etablissement("plafond")
    demande = _demande(matieres, classes, profs, limite_secondes=999_999)
    tache_id = client.post(f"/emplois-du-temps/{edt}/generer", json=demande,
                           headers=entetes).json()["id"]
    import json as _json
    from app.database import SessionLocal
    from app.models.task import TacheGeneration
    db = SessionLocal()
    try:
        stockee = _json.loads(db.get(TacheGeneration, tache_id).requete)
        assert stockee["limite_secondes"] == settings.LIMITE_SECONDES_MAX
    finally:
        db.close()
    client.delete(f"/emplois-du-temps/{edt}/taches/{tache_id}", headers=entetes)
    _attendre(entetes, edt, tache_id)


# ══════════════════════════════════════════════════════════════════
#  Jeu de démonstration
# ══════════════════════════════════════════════════════════════════

def test_demonstration_apercu_ne_cree_rien():
    entetes, *_ = _etablissement("demo-apercu")
    apercu = client.get("/demonstration/", headers=entetes).json()
    assert apercu["classes"] > 0 and apercu["professeurs"] > 0
    # L'établissement de test contient déjà des ressources.
    assert apercu["etablissement_deja_peuple"] is True

    vierge = client.post("/auth/inscrire", json={
        "ecole": {"nom": "Vierge", "email": "vierge@test.dz"},
        "admin": {"nom": "A", "prenom": "B", "email": "vierge-admin@test.dz",
                  "mot_de_passe": "motdepasse-solide"}})
    h = {"Authorization": f"Bearer {vierge.json()['access_token']}"}
    assert client.get("/demonstration/", headers=h).json()[
        "etablissement_deja_peuple"] is False
    # L'aperçu n'a rien écrit.
    assert client.get("/matieres/", headers=h).json() == []


def test_demonstration_refuse_decraser_sans_confirmation():
    """Effacer le travail d'un responsable sans le lui demander serait pire
    que de ne rien faire."""
    entetes, matieres, classes, profs, edt = _etablissement("demo-ecrase")
    avant = len(client.get("/matieres/", headers=entetes).json())

    r = client.post("/demonstration/charger", json={}, headers=entetes)
    assert r.status_code == 409
    assert len(client.get("/matieres/", headers=entetes).json()) == avant

    r = client.post("/demonstration/charger", json={"remplacer": True},
                    headers=entetes)
    assert r.status_code == 201, r.text
    apres = client.get("/matieres/", headers=entetes).json()
    assert len(apres) == r.json()["matieres"]
    # Les emplois du temps de l'établissement ont été remplacés eux aussi.
    edts = client.get("/emplois-du-temps/", headers=entetes).json()
    assert [e["id"] for e in edts] == [r.json()["emploi_du_temps_id"]]


def test_demonstration_produit_un_emploi_du_temps_complet():
    """
    Le jeu doit se générer sans aucune intervention : programme,
    réglages et fenêtres sont installés avec lui. C'est tout son intérêt.
    """
    r = client.post("/auth/inscrire", json={
        "ecole": {"nom": "Demo", "email": "demo-gen@test.dz"},
        "admin": {"nom": "A", "prenom": "B", "email": "demo-gen-a@test.dz",
                  "mot_de_passe": "motdepasse-solide"}})
    entetes = {"Authorization": f"Bearer {r.json()['access_token']}"}

    installation = client.post("/demonstration/charger", json={},
                               headers=entetes)
    assert installation.status_code == 201, installation.text
    resume = installation.json()
    edt = resume["emploi_du_temps_id"]

    programme = client.get("/programme/", headers=entetes).json()
    assert len(programme) == resume["lignes_programme"]
    assert sum(c["heures_par_semaine"] for c in programme) == \
        resume["lecons_a_placer"]
    # Les dédoublements en demi-groupes sont repris tels quels.
    assert {l["couplage_id"] for l in programme if l["couplage_id"]}

    reglages = client.get("/parametres", headers=entetes).json()
    assert reglages["grille"]["fermetures"], "le mardi après-midi doit être fermé"
    fenetres = client.get("/fenetres-pedagogiques", headers=entetes).json()
    assert len(fenetres) == resume["fenetres_pedagogiques"]

    # Rien d'autre que le temps de calcul n'est envoyé.
    demande = {"limite_secondes": 120}
    diagnostic = client.post(f"/emplois-du-temps/{edt}/diagnostic",
                             json=demande, headers=entetes).json()
    assert diagnostic["realisable"], diagnostic["erreurs"]

    tache = client.post(f"/emplois-du-temps/{edt}/generer", json=demande,
                        headers=entetes)
    assert tache.status_code == 202
    etat = _attendre(entetes, edt, tache.json()["id"], delai=240)
    assert etat["statut"] == "terminee", etat["message"]
    assert etat["resultat"]["lecons_planifiees"] == resume["lecons_a_placer"]
    assert etat["resultat"]["qualite"]["trous_classes"] == 0

    lecons = client.get(f"/emplois-du-temps/{edt}/lecons", headers=entetes).json()

    # La fermeture enregistrée ne porte aucune leçon.
    jours = reglages["grille"]["jours"]
    horaires = [h[0] for h in reglages["grille"]["horaires"]]
    fermes = {(jours[j], horaires[s])
              for j, seances in reglages["grille"]["fermetures"] for s in seances}
    assert not [l for l in lecons if (l["jour"], l["heure_debut"]) in fermes]

    # Aucune matière ne tombe dans sa fenêtre d'inspection.
    interdits = {(f["matiere_id"], jours[f["index_jour"]], horaires[s])
                 for f in fenetres for s in f["seances_bloquees"]}
    for l in lecons:
        assert (l["matiere_id"], l["jour"], l["heure_debut"]) not in interdits


# ══════════════════════════════════════════════════════════════════
#  Réglages persistés, fenêtres pédagogiques et fouj
# ══════════════════════════════════════════════════════════════════

def test_parametres_valeurs_par_defaut_sans_reglage():
    entetes, *_ = _etablissement("param-defaut")
    p = client.get("/parametres", headers=entetes).json()
    assert p["grille"]["jours"][0] == "Dimanche"
    assert len(p["grille"]["horaires"]) == 7
    assert p["grille"]["fermetures"] == []
    assert p["ponderations"]["penalites_seance"]["6"] == 150
    assert p["ponderations"]["penalites_heures_par_jour"] == {"5": 15, "6": 45}


def test_parametres_enregistres_et_relus():
    entetes, *_ = _etablissement("param-ecrit")
    grille = client.get("/parametres", headers=entetes).json()["grille"]
    grille["fermetures"] = [[2, [4, 5, 6]]]

    poids = client.get("/parametres", headers=entetes).json()["ponderations"]
    poids["penalites_heures_par_jour"] = {"4": 5, "5": 20, "6": 60}

    r = client.put("/parametres", json={
        "grille": grille,
        "ponderations": poids,
        "presence_minimale": {"matin": 3},
        "limite_secondes": 45,
    }, headers=entetes)
    assert r.status_code == 200, r.text

    relu = client.get("/parametres", headers=entetes).json()
    assert relu["grille"]["fermetures"] == [[2, [4, 5, 6]]]
    assert relu["presence_minimale"] == {"matin": 3}
    assert relu["limite_secondes"] == 45
    # Les seuils sont indexés par des entiers côté solveur et par du texte
    # en JSON : la relecture doit rendre exactement ce qui a été envoyé.
    assert relu["ponderations"]["penalites_heures_par_jour"] == {
        "4": 5, "5": 20, "6": 60}


def test_grille_incoherente_est_refusee():
    """Une grille invalide rendrait toute génération impossible sans que
    la cause soit visible."""
    entetes, *_ = _etablissement("param-invalide")
    grille = client.get("/parametres", headers=entetes).json()["grille"]

    # Une séance qui n'appartient à aucune demi-journée.
    casse = dict(grille, shifts=[["matin", [0, 1, 2, 3]], ["apres-midi", [4, 5]]])
    r = client.put("/parametres", json={"grille": casse}, headers=entetes)
    assert r.status_code == 422 and "demi-journée" in r.json()["detail"]

    # Aucun jour travaillé.
    r = client.put("/parametres", json={"grille": dict(grille, jours=[])},
                   headers=entetes)
    assert r.status_code == 422


def test_generation_reprend_les_reglages_enregistres():
    """Le client n'envoie que les cours ; le reste vient des réglages."""
    entetes, matieres, classes, profs, edt = _etablissement("param-generation")
    grille = client.get("/parametres", headers=entetes).json()["grille"]
    grille["fermetures"] = [[2, [4, 5, 6]]]          # mardi après-midi fermé
    client.put("/parametres", json={"grille": grille}, headers=entetes)

    # Fenêtre : pas de sport le dimanche matin.
    fenetre = client.post("/fenetres-pedagogiques", json={
        "matiere_id": matieres["Sport"], "index_jour": 0,
        "seances_bloquees": [0, 1, 2, 3], "libelle": "Inspection EPS"},
        headers=entetes)
    assert fenetre.status_code == 201, fenetre.text

    demande = {"cours_requis": _demande(matieres, classes, profs)["cours_requis"],
               "limite_secondes": 30}
    tache = client.post(f"/emplois-du-temps/{edt}/generer", json=demande,
                        headers=entetes)
    assert tache.status_code == 202, tache.text
    etat = _attendre(entetes, edt, tache.json()["id"], delai=90)
    assert etat["statut"] == "terminee", etat["message"]

    lecons = client.get(f"/emplois-du-temps/{edt}/lecons", headers=entetes).json()
    # La fermeture enregistrée a été appliquée.
    assert not [l for l in lecons if l["jour"] == "Mardi" and l["heure_debut"] >= "13:00"]
    # La fenêtre pédagogique aussi.
    sport = [l for l in lecons if l["matiere_nom"] == "Sport"]
    assert sport and not [l for l in sport
                          if l["jour"] == "Dimanche" and l["heure_debut"] < "12:00"]


def test_fenetre_en_double_est_refusee():
    entetes, matieres, *_ = _etablissement("fenetre-double")
    corps = {"matiere_id": matieres["Maths"], "index_jour": 1,
             "seances_bloquees": [0, 1]}
    assert client.post("/fenetres-pedagogiques", json=corps,
                       headers=entetes).status_code == 201
    r = client.post("/fenetres-pedagogiques", json=corps, headers=entetes)
    assert r.status_code == 409

    liste = client.get("/fenetres-pedagogiques", headers=entetes).json()
    assert len(liste) == 1 and liste[0]["seances_bloquees"] == [0, 1]
    client.delete(f"/fenetres-pedagogiques/{liste[0]['id']}", headers=entetes)
    assert client.get("/fenetres-pedagogiques", headers=entetes).json() == []


def test_fouj_deux_demi_groupes_au_meme_creneau():
    """Les deux moitiés d'une classe suivent deux cours simultanés, dans
    deux salles et avec deux enseignants distincts."""
    entetes, matieres, classes, profs, edt = _etablissement("fouj")
    classe = classes[0]
    couplage = "fouj-maths-francais"
    cours = [
        {"classe_id": classe, "matiere_id": matieres["Maths"],
         "professeur_id": profs["Maths"][0], "heures_par_semaine": 2,
         "nb_seances_doubles": 0, "max_heures_par_jour": 1,
         "couplage_id": couplage, "groupe": "G1"},
        {"classe_id": classe, "matiere_id": matieres["Français"],
         "professeur_id": profs["Français"][0], "heures_par_semaine": 2,
         "nb_seances_doubles": 0, "max_heures_par_jour": 1,
         "couplage_id": couplage, "groupe": "G2"},
        {"classe_id": classe, "matiere_id": matieres["Arabe"],
         "professeur_id": profs["Arabe"][0], "heures_par_semaine": 3},
    ]
    tache = client.post(f"/emplois-du-temps/{edt}/generer",
                        json={"cours_requis": cours, "limite_secondes": 30},
                        headers=entetes)
    assert tache.status_code == 202, tache.text
    etat = _attendre(entetes, edt, tache.json()["id"], delai=90)
    assert etat["statut"] == "terminee", etat["message"]

    lecons = client.get(f"/emplois-du-temps/{edt}/lecons", headers=entetes).json()
    maths = {(l["jour"], l["heure_debut"]) for l in lecons
             if l["matiere_nom"] == "Maths"}
    francais = {(l["jour"], l["heure_debut"]) for l in lecons
                if l["matiere_nom"] == "Français"}
    assert maths and maths == francais, (maths, francais)

    # Même créneau, mais jamais la même salle.
    for jour, heure in maths:
        salles = {l["salle_id"] for l in lecons
                  if (l["jour"], l["heure_debut"]) == (jour, heure)}
        assert len(salles) == 2, f"{jour} {heure} : {salles}"


# ══════════════════════════════════════════════════════════════════
#  Programme annuel
# ══════════════════════════════════════════════════════════════════

def test_programme_persiste_et_alimente_la_generation():
    """Le client ne doit plus envoyer les cours : ils sont en base."""
    entetes, matieres, classes, profs, edt = _etablissement("prog-persiste")
    demande = _demande(matieres, classes, profs)

    for ligne in demande["cours_requis"]:
        r = client.post("/programme/", json=ligne, headers=entetes)
        assert r.status_code == 201, r.text

    enregistre = client.get("/programme/", headers=entetes).json()
    assert len(enregistre) == len(demande["cours_requis"])
    assert enregistre[0]["classe_nom"] and enregistre[0]["professeur_nom"]

    # Génération sans « cours_requis » : le programme enregistré est repris.
    tache = client.post(f"/emplois-du-temps/{edt}/generer",
                        json={"limite_secondes": 25}, headers=entetes)
    assert tache.status_code == 202, tache.text
    etat = _attendre(entetes, edt, tache.json()["id"], delai=90)
    assert etat["statut"] == "terminee", etat["message"]
    attendu = sum(c["heures_par_semaine"] for c in demande["cours_requis"])
    assert etat["resultat"]["lecons_planifiees"] == attendu


def test_generation_sans_programme_est_refusee():
    entetes, _, _, _, edt = _etablissement("prog-vide")
    client.delete("/programme/", headers=entetes)
    r = client.post(f"/emplois-du-temps/{edt}/generer",
                    json={"limite_secondes": 20}, headers=entetes)
    assert r.status_code == 400
    assert "programme" in str(r.json()["detail"]).lower()


def test_import_csv_aller_retour_fidele():
    entetes, matieres, classes, profs, edt = _etablissement("prog-csv")
    for ligne in _demande(matieres, classes, profs)["cours_requis"]:
        client.post("/programme/", json=ligne, headers=entetes)

    exporte = client.get("/programme/export", headers=entetes).text
    avant = client.get("/programme/", headers=entetes).json()

    r = client.post("/programme/importer",
                    files={"fichier": ("p.csv", exporte.encode(), "text/csv")},
                    headers=entetes)
    assert r.status_code == 200, r.text
    apres = client.get("/programme/", headers=entetes).json()

    cle = lambda l: (l["classe_nom"], l["matiere_nom"], l["professeur_nom"],
                     l["heures_par_semaine"], l["nb_seances_doubles"])
    assert sorted(map(cle, avant)) == sorted(map(cle, apres))


def test_import_fautif_n_ecrit_rien():
    """Un programme à moitié chargé serait pire qu'un import refusé."""
    entetes, matieres, classes, profs, edt = _etablissement("prog-fautif")
    for ligne in _demande(matieres, classes, profs)["cours_requis"][:3]:
        client.post("/programme/", json=ligne, headers=entetes)
    avant = client.get("/programme/", headers=entetes).json()

    mauvais = (
        "classe,matiere,professeur,heures\n"
        f"{classes[0] and '1AM1'},Maths,PMaths0,4\n"
        "ClasseInconnue,Maths,PMaths0,4\n"
        "1AM1,MatiereInconnue,PMaths0,4\n"
        "1AM1,Maths,PMaths0,zero\n"
    )
    r = client.post("/programme/importer",
                    files={"fichier": ("m.csv", mauvais.encode(), "text/csv")},
                    headers=entetes)
    assert r.status_code == 422
    detail = r.json()["detail"]
    assert detail["erreurs_totales"] >= 3
    assert any("inconnue" in m for m in detail["erreurs"])
    # Rien n'a bougé.
    assert client.get("/programme/", headers=entetes).json() == avant


def test_import_csv_point_virgule_et_entete_anglais():
    """Les tableurs francophones exportent en point-virgule."""
    entetes, matieres, classes, profs, edt = _etablissement("prog-sep")
    contenu = ("Class;Subject;Teacher;Hours;blocs_2h\n"
               "1AM1;Maths;PMaths0;4;1\n"
               "1AM1;Arabe;PArabe0;4;0\n")
    r = client.post("/programme/importer",
                    files={"fichier": ("s.csv", contenu.encode(), "text/csv")},
                    headers=entetes)
    assert r.status_code == 200, r.text
    lignes = client.get("/programme/", headers=entetes).json()
    assert len(lignes) == 2
    assert {l["matiere_nom"] for l in lignes} == {"Maths", "Arabe"}


def test_import_fouj_par_paires_seulement():
    entetes, matieres, classes, profs, edt = _etablissement("prog-fouj-csv")
    seul = ("classe,matiere,professeur,heures,fouj,groupe\n"
            "1AM1,Maths,PMaths0,2,DEDOUBLE,G1\n")
    r = client.post("/programme/importer",
                    files={"fichier": ("f.csv", seul.encode(), "text/csv")},
                    headers=entetes)
    assert r.status_code == 422
    assert "2" in str(r.json()["detail"]["erreurs"])

    paire = ("classe,matiere,professeur,heures,fouj,groupe\n"
             "1AM1,Maths,PMaths0,2,DEDOUBLE,G1\n"
             "1AM1,Arabe,PArabe0,2,DEDOUBLE,G2\n")
    r = client.post("/programme/importer",
                    files={"fichier": ("f.csv", paire.encode(), "text/csv")},
                    headers=entetes)
    assert r.status_code == 200, r.text
    lignes = client.get("/programme/", headers=entetes).json()
    assert len({l["couplage_id"] for l in lignes}) == 1
    assert {l["groupe"] for l in lignes} == {"G1", "G2"}

    # Supprimer une moitié emporte l'autre : une moitié seule est invalide.
    client.delete(f"/programme/{lignes[0]['id']}", headers=entetes)
    assert client.get("/programme/", headers=entetes).json() == []


def test_programme_cloisonne_entre_etablissements():
    entetes_a, matieres, classes, profs, _ = _etablissement("prog-tenant-a")
    entetes_b, *_ = _etablissement("prog-tenant-b")
    client.post("/programme/", json=_demande(matieres, classes, profs)
                ["cours_requis"][0], headers=entetes_a)
    assert client.get("/programme/", headers=entetes_a).json()
    assert client.get("/programme/", headers=entetes_b).json() == []


if __name__ == "__main__":
    echecs = 0
    for nom, fonction in sorted(globals().items()):
        if not nom.startswith("test_") or not callable(fonction):
            continue
        depart = time.perf_counter()
        try:
            fonction()
            print(f"  ✓ {nom}  ({time.perf_counter() - depart:.1f}s)")
        except Exception as e:
            echecs += 1
            import traceback
            print(f"  ✗ {nom}\n      {e or type(e).__name__}")
            print("      " + traceback.format_exc().replace("\n", "\n      ")[:900])
    if BASE.exists():
        BASE.unlink()
    print(f"\n{'Tous les tests passent.' if not echecs else f'{echecs} échec(s).'}")
    sys.exit(1 if echecs else 0)
